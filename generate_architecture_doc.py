"""generate_architecture_doc.py
Generates the technical architecture & functionality Word document for the
Treasury & Commodity Counterparty Intelligence platform.
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

DELOITTE_GREEN = RGBColor(0x86, 0xBC, 0x25)
BLACK = RGBColor(0x00, 0x00, 0x00)
GRAY = RGBColor(0x4B, 0x55, 0x63)

doc = Document()

# Base style
normal = doc.styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(11)

def h1(text):
    p = doc.add_heading(level=1)
    r = p.add_run(text)
    r.font.color.rgb = BLACK
    r.font.name = "Aptos"
    return p

def h2(text):
    p = doc.add_heading(level=2)
    r = p.add_run(text)
    r.font.color.rgb = DELOITTE_GREEN
    r.font.name = "Aptos"
    return p

def para(text, bold=False, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    return p

def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p

def num(text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p

def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, htext in enumerate(headers):
        hdr[i].paragraphs[0].add_run(htext).bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    doc.add_paragraph()
    return t

# ── Title page ──────────────────────────────────────────────────────────────
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("Treasury & Commodity Counterparty Intelligence")
r.bold = True
r.font.size = Pt(26)
r.font.color.rgb = BLACK
r.font.name = "Aptos"

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
rs = sub.add_run("Technical Architecture & Code Functionality — Google Cloud Serverless Platform")
rs.font.size = Pt(13)
rs.font.color.rgb = DELOITTE_GREEN
rs.font.name = "Aptos"

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run("Deloitte · Counterparty Risk Assessment Platform").italic = True

doc.add_paragraph()

# ── 1. Executive Overview ─────────────────────────────────────────────────────
h1("1. Executive Overview")
para(
    "The platform is an event-driven, serverless application on Google Cloud Platform (GCP) "
    "that ingests counterparty financial statements, stores structured data in BigQuery, and "
    "applies a set of specialist large-language-model (LLM) agents to produce market, credit, "
    "liquidity, and cashflow risk assessments. The user interacts through a single-page web "
    "dashboard hosted on Cloud Run, which serves both the UI and the JSON APIs. All LLM calls "
    "use Google Gemini as the primary engine with Anthropic Claude as an automatic fallback, "
    "and live external context is fetched from Tavily web search, Yahoo Finance, and SEC EDGAR."
)

# ── 2. GCP Architecture ───────────────────────────────────────────────────────
h1("2. Google Cloud Setup")

h2("2.1 Core Services")
table(
    ["Service", "Role in the Platform"],
    [
        ["Cloud Run", "Hosts two containers: 'treasury-ingestor' (background ingestion) and 'treasury-dashboard' (web UI + APIs). Deployed privately (no unauthenticated access)."],
        ["BigQuery", "Analytical store. Dataset 'treasury_analytics' holds counterparties, agent_memos, and deposits tables."],
        ["Cloud Storage", "Landing bucket for uploaded financial documents; emits object-create notifications."],
        ["Pub/Sub", "Topic 'treasury-financials' receives GCS notifications and push-delivers them to the ingestor '/ingest' endpoint."],
        ["Artifact Registry", "Stores the Docker images ('treasury-repo') built by Cloud Build."],
        ["Cloud Build", "Builds and pushes container images from source."],
        ["HTTPS Load Balancer + IAP", "Global external load balancer with a serverless NEG and Identity-Aware Proxy for corporate SSO access control."],
    ],
)

h2("2.2 Event-Driven Ingestion Pipeline")
num("A user (or batch job) uploads a CSV/XLSX financial statement to the Cloud Storage bucket.")
num("Cloud Storage emits an object-finalize notification to the Pub/Sub topic 'treasury-financials'.")
num("Pub/Sub push-delivers the message (with an OIDC token) to the ingestor service '/ingest' endpoint.")
num("The ingestor downloads the blob, parses it with the OCR/parsing module, and streams normalized rows into the BigQuery 'counterparties' table.")
num("The dashboard reads from BigQuery to render KPIs, tables, and to ground the LLM agents.")

h2("2.3 Security & Configuration")
bullet("Authentication: Cloud Run services run privately; public access is mediated by the HTTPS Load Balancer with Identity-Aware Proxy (IAP). The app also verifies the X-Goog-IAP-JWT-Assertion header.")
bullet("Secrets: API keys are injected as Cloud Run environment variables — GEMINI_API_KEY, ANTHROPIC_API_KEY, TAVILY_API_KEY — plus BUCKET and BQ_DATASET. They are never committed to source control (.env is git-ignored).")
bullet("Least privilege: a dedicated invoker service account allows Pub/Sub to call the private Cloud Run ingestion endpoint.")

# ── 3. Application Architecture ───────────────────────────────────────────────
h1("3. Application & Code Functionality")

h2("3.1 Module Map")
table(
    ["Module", "Responsibility"],
    [
        ["dashboard.py", "Flask app: serves dashboard.html and all JSON APIs (counterparties, stats, upload, chat, agent endpoints, liquidity-model, memos, Pub/Sub ingest)."],
        ["agent_base.py", "Shared LLM layer: call_llm (Gemini primary, Claude fallback) with reasoning trace logging; web_search (Tavily); standardized reporting helpers; BigQuery memo persistence."],
        ["market_agent.py", "Market Risk Agent — commodity price sensitivity, ownership/IR, recommended notional exposure."],
        ["credit_agent.py", "Credit Risk Agent — creditworthiness, leverage, credit limit and payment terms."],
        ["liquidity_agent.py", "Liquidity & Settlement Agent — settlement capability on T+5 to T+10 terms."],
        ["liquidity_model_agent.py", "Liquidity Modelling Agent — 3-week (21-day) cashflow projection from uploaded operational files."],
        ["bigquery_helper.py", "All BigQuery reads/writes and LLM context assembly."],
        ["market_data_helper.py", "Live market context: Yahoo Finance quotes/ownership and SEC EDGAR filings presence."],
        ["risk_scorer.py", "Deterministic 0–100 quantitative risk score used alongside the LLM narrative."],
        ["ocr_simulator.py / claude_ocr_enhance.py", "Financial document parsing with a Claude vision fallback."],
    ],
)

h2("3.2 LLM Orchestration & Fallback")
para(
    "Every LLM interaction flows through call_llm() in agent_base.py. It first attempts Google "
    "Gemini; if Gemini is unavailable or errors, it automatically falls back to Anthropic Claude. "
    "The function logs a preview of each prompt and the resulting reasoning output, giving full "
    "traceability of how the models reason about each counterparty. A trace_label parameter tags "
    "each call (MARKET, CREDIT, LIQUIDITY, LIQ_MODEL) so the logs are easy to follow."
)

h2("3.3 Standardized, Actionable Reporting")
para(
    "To keep output consistent and decision-ready, all agents share a STANDARD_REPORT_INSTRUCTION "
    "and a summarize_to_sentences() safety net. Each memo is constrained to a maximum of six "
    "concise, quantitative sentences ending in an actionable conclusion on exposure or credit "
    "capacity. Each agent also returns a structured verdict (risk level plus an exposure / credit "
    "limit / settlement-terms field) and a deterministic risk score, which are persisted to the "
    "agent_memos table for audit."
)

h2("3.4 Risk Agents")
bullet("Market Risk Agent: assesses revenue quality, commodity price sensitivity, ownership structure and investor relations, and recommends a maximum notional trading exposure.")
bullet("Credit Risk Agent: evaluates credit rating, leverage and debt serviceability, then recommends an unsecured credit limit and payment terms.")
bullet("Liquidity & Settlement Agent: judges short-term solvency and settlement reliability on standard LNG terms, recommending a settlement structure.")
bullet("Orchestrator: runs the three agents in parallel, computes a weighted aggregate risk (market 40%, credit 35%, liquidity 25%), and synthesizes an executive recommendation (approve / approve with conditions / decline).")

h2("3.5 Market Data Assistant (Chat)")
para(
    "The Market Data Assistant endpoint (/api/chat) grounds the LLM in three layers: internal "
    "BigQuery data (financial snapshots, deposits, prior agent memos), live market and ownership "
    "context (Yahoo Finance and SEC EDGAR), and live Tavily web search focused on ownership and "
    "investor relations. Responses follow the same six-sentence standardized format."
)

h2("3.6 Liquidity Modelling Agent (3-Week Cashflow)")
para(
    "The Liquidity Modelling Agent accepts one or more operational files — account receivables, "
    "payables, taxes, maturity of financial swaps, and margin requirements — and projects the "
    "next 21 days of cashflow. A deterministic engine parses each file (flexible, case-insensitive "
    "column mapping), classifies every line as an inflow or outflow, and computes a daily running "
    "balance plus a three-week summary and a net-by-category breakdown. The LLM is used only for a "
    "concise treasury commentary that flags funding gaps and recommends an action. Optionally it "
    "also pulls deposits on file from BigQuery as additional inflows. The dashboard renders the "
    "result as KPIs, an inline SVG chart (daily balance line with net-flow bars), and tabular output."
)

# ── 4. API Endpoints ──────────────────────────────────────────────────────────
h1("4. API Endpoints")
table(
    ["Method & Path", "Description"],
    [
        ["GET /", "Serves the dashboard single-page UI."],
        ["GET /api/counterparties", "Counterparty records (optional search/sector filters)."],
        ["GET /api/stats", "Portfolio KPI summary."],
        ["POST /api/upload", "Local file ingestion into BigQuery."],
        ["POST /api/chat", "Market Data Assistant Q&A (BigQuery + market + web grounded)."],
        ["POST /api/agent/market|credit|liquidity", "Run an individual specialist risk agent."],
        ["POST /api/agent/orchestrate", "Run all three agents and synthesize an executive summary."],
        ["POST /api/liquidity-model", "3-week cashflow projection from uploaded operational file(s)."],
        ["GET /api/memos", "Retrieve stored agent assessments."],
        ["POST /ingest", "Pub/Sub push handler for background ingestion."],
    ],
)

# ── 5. Data Model ─────────────────────────────────────────────────────────────
h1("5. BigQuery Data Model")
h2("5.1 counterparties")
para("Financial snapshots per counterparty: company_name, country, sector, credit_rating, "
     "period_year, revenue_usd_m, ebitda_usd_m, ebitda_margin_pct, net_income_usd_m, "
     "total_assets_usd_m, total_debt_usd_m, debt_to_equity, current_ratio, document_name, upload_date.")
h2("5.2 agent_memos")
para("Agent outputs: id, counterparty_name, agent_type, risk_level, memo, exposure_proposal, "
     "created_at, plus extended JSON fields (credit_limit, payment_terms, settlement_terms, "
     "risk_score, risk_score_breakdown).")
h2("5.3 deposits")
para("Deposit records used for liquidity context and as inflows in the cashflow model.")

# ── 6. Deployment ─────────────────────────────────────────────────────────────
h1("6. Deployment Summary")
num("Enable required APIs (Cloud Run, Artifact Registry, Cloud Build, Compute, IAP).")
num("Create the BigQuery dataset and tables (idempotent).")
num("Create the Cloud Storage bucket, Pub/Sub topic, and GCS notification.")
num("Build the container image with Cloud Build and push to Artifact Registry.")
num("Deploy 'treasury-ingestor' and 'treasury-dashboard' to Cloud Run (private).")
num("Provision the HTTPS Load Balancer, serverless NEG, and enable IAP for corporate access.")
num("Point the Pub/Sub push subscription at the ingestion endpoint using the invoker service account.")

doc.add_paragraph()
footer = doc.add_paragraph()
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = footer.add_run("Deloitte · Treasury & Commodity Platform · Powered by Google Cloud")
fr.italic = True
fr.font.color.rgb = GRAY

out = "Treasury_Platform_Architecture.docx"
doc.save(out)
print(f"Saved {out}")
